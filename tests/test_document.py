"""tests/test_document.py

Purpose
-------
Offline unittest coverage for document.py: section-specific header
extraction, .docx content/exclusion rules, the bounded formatting-preset
loop, and the render-only CLI's validation gate. All synthetic data -
never resume_profile.md.

TestRenderAndVerifyMocked patches document.convert_docx_to_pdf and
document.count_pdf_pages directly: it checks render_and_verify()'s own
control flow (bounded attempts, which status each page count produces),
not real LibreOffice behavior. The actual DOCX-to-PDF rendering smoke
test (real LibreOffice, real page count) is run separately, outside
this file, exactly because it is not something a unit test should
depend on.

Run with: .venv/Scripts/python.exe -m unittest discover -s tests -t . -v
"""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document as DocxDocument

from retrieval import Chunk
from decision import Bullet, DraftResult, EvidenceDecision, ResumeEntry, ResumeSection, decision_to_dict
import document
from document import (
    DocumentGenerationError,
    FORMATTING_PRESETS,
    MIN_BODY_FONT_PT,
    MIN_MARGIN_IN,
    LibreOfficeConversionError,
    build_docx,
    extract_education_header,
    extract_entry_header,
    extract_job_header,
    load_saved_draft,
    render_and_verify,
)


def make_chunk(chunk_id: str, section: str, title: str, text: str) -> Chunk:
    return Chunk(chunk_id=chunk_id, section=section, title=title, text=text)


GOOD_JOB_CHUNK = make_chunk(
    "job-1",
    "job",
    "Acme Corp — Data Analyst",
    "### Acme Corp — Data Analyst\n\n- Location: Springfield, IL.\n- Dates: 2020–2023; full-time.\n- Did useful analyst things.",
)
REORDERED_JOB_CHUNK = make_chunk(
    "job-2",
    "job",
    "Beta Inc — Engineer",
    "### Beta Inc — Engineer\n\n- Dates: 2018–2020; part-time.\n- Location: Metropolis, NY.\n- Did engineering things.",
)
BAD_JOB_CHUNK = make_chunk(
    "job-3",
    "job",
    "Gamma LLC — Tech",
    "### Gamma LLC — Tech\n\n- Location: Nowhere, TX.\n- Did some tech stuff.",
)
GOOD_EDU_CHUNK = make_chunk(
    "edu-1",
    "education_training",
    "B.S. in Widgetry",
    "### B.S. in Widgetry\n\n- Springfield University; completed 2015; GPA 3.9.\n- Took various widget courses.",
)
BAD_EDU_CHUNK = make_chunk(
    "edu-2",
    "education_training",
    "Certificate in Things",
    "### Certificate in Things\n\n- Completed a short course in 2021 at Acme Academy.\n- Learned things.",
)
PROJECT_CHUNK = make_chunk(
    "proj-1", "project", "Widget Tracker App", "### Widget Tracker App\n\n- Built a tracker.\n- Tools: Python."
)
SKILL_CHUNK = make_chunk(
    "skill-1", "skill_category", "Programming languages", "- **Programming languages:** Python, SQL."
)

ALL_CHUNKS = [GOOD_JOB_CHUNK, GOOD_EDU_CHUNK, PROJECT_CHUNK, SKILL_CHUNK]

SECRET_REASON = "SECRET_REASON_SHOULD_NOT_APPEAR_IN_DOCX"

GOOD_CONTACT = {"Name": "Jordan Example", "Phone": "+1 555-0100", "Email": "jordan@example.com"}

GOOD_DECISION = DraftResult(
    professional_summary="Experienced analyst with a strong technical foundation.",
    summary_evidence_ids=["job-1"],
    resume_sections=[
        ResumeSection(
            section="skill_category",
            entries=[ResumeEntry(source_chunk_id="skill-1", bullets=[Bullet(text="Python, SQL.", evidence_ids=["skill-1"])])],
        ),
        ResumeSection(
            section="job",
            entries=[
                ResumeEntry(
                    source_chunk_id="job-1",
                    bullets=[Bullet(text="Did useful analyst things.", evidence_ids=["job-1"])],
                )
            ],
        ),
        ResumeSection(
            section="education_training",
            entries=[
                ResumeEntry(
                    source_chunk_id="edu-1",
                    bullets=[Bullet(text="Took various widget courses.", evidence_ids=["edu-1"])],
                )
            ],
        ),
        ResumeSection(
            section="project",
            entries=[ResumeEntry(source_chunk_id="proj-1", bullets=[Bullet(text="Built a tracker.", evidence_ids=["proj-1"])])],
        ),
    ],
    selected_evidence=[EvidenceDecision(chunk_id=cid, reason=SECRET_REASON) for cid in ["job-1", "edu-1", "proj-1", "skill-1"]],
    omitted_evidence=[],
)


class TestExtractEducationHeader(unittest.TestCase):
    def test_extracts_institution_and_date(self):
        result = extract_education_header(GOOD_EDU_CHUNK)
        self.assertTrue(result.ok)
        self.assertEqual(result.institution, "Springfield University")
        self.assertEqual(result.dates, "completed 2015")

    def test_reports_failure_for_unmatched_first_line(self):
        result = extract_education_header(BAD_EDU_CHUNK)
        self.assertFalse(result.ok)
        self.assertIn("edu-2", result.reason)


class TestExtractJobHeader(unittest.TestCase):
    def test_extracts_location_and_dates(self):
        result = extract_job_header(GOOD_JOB_CHUNK)
        self.assertTrue(result.ok)
        self.assertEqual(result.location, "Springfield, IL.")
        self.assertEqual(result.dates, "2020–2023; full-time.")

    def test_tolerates_reordered_location_and_dates_lines(self):
        result = extract_job_header(REORDERED_JOB_CHUNK)
        self.assertTrue(result.ok)
        self.assertEqual(result.location, "Metropolis, NY.")
        self.assertEqual(result.dates, "2018–2020; part-time.")

    def test_reports_failure_when_dates_line_missing(self):
        result = extract_job_header(BAD_JOB_CHUNK)
        self.assertFalse(result.ok)
        self.assertIn("job-3", result.reason)
        self.assertIn("Dates:", result.reason)


class TestExtractEntryHeaderDispatch(unittest.TestCase):
    def test_project_and_skill_have_no_header_to_extract(self):
        self.assertIsNone(extract_entry_header(PROJECT_CHUNK))
        self.assertIsNone(extract_entry_header(SKILL_CHUNK))


class TestFormattingPresets(unittest.TestCase):
    def test_presets_never_violate_the_readability_floors(self):
        self.assertLessEqual(len(FORMATTING_PRESETS), 3)
        for preset in FORMATTING_PRESETS:
            self.assertGreaterEqual(preset.body_font_pt, MIN_BODY_FONT_PT)
            self.assertGreaterEqual(preset.margin_in, MIN_MARGIN_IN)


class TestBuildDocx(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.docx_path = Path(self._tmp.name) / "resume.docx"

    def test_builds_expected_content_and_excludes_evidence_trace(self):
        build_docx(GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, FORMATTING_PRESETS[0])

        doc = DocxDocument(str(self.docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)

        self.assertIn("Jordan Example", full_text)
        self.assertIn(GOOD_DECISION.professional_summary, full_text)
        self.assertIn("Professional Experience", full_text)
        self.assertIn("Education and Training", full_text)
        self.assertIn("Acme Corp — Data Analyst", full_text)
        self.assertIn("Springfield University", full_text)  # exact extracted institution
        self.assertIn("completed 2015", full_text)  # exact extracted date phrase
        self.assertIn("Did useful analyst things.", full_text)

        # Evidence IDs / selection reasons must never appear in the document.
        self.assertNotIn(SECRET_REASON, full_text)
        self.assertNotIn("job-1", full_text)

    def test_sets_letter_page_size_calibri_font_and_bullet_size(self):
        preset = FORMATTING_PRESETS[1]  # "compact": 10.5pt
        build_docx(GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, preset)

        doc = DocxDocument(str(self.docx_path))
        section_fmt = doc.sections[0]
        self.assertAlmostEqual(section_fmt.page_width.inches, 8.5)
        self.assertAlmostEqual(section_fmt.page_height.inches, 11.0)

        normal = doc.styles["Normal"]
        self.assertEqual(normal.font.name, "Calibri")
        self.assertAlmostEqual(normal.font.size.pt, preset.body_font_pt)

        list_bullet = doc.styles["List Bullet"]
        self.assertEqual(list_bullet.font.name, "Calibri")
        self.assertAlmostEqual(list_bullet.font.size.pt, preset.body_font_pt)

    def test_major_headings_are_12pt_bold_black_with_spacing_and_kept_with_next(self):
        preset = FORMATTING_PRESETS[0]
        build_docx(GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, preset)

        doc = DocxDocument(str(self.docx_path))
        heading_texts = {"Professional Summary", "Professional Experience", "Education and Training", "Projects", "Skills"}
        headings = [p for p in doc.paragraphs if p.text in heading_texts]
        self.assertTrue(headings)

        for heading in headings:
            run = heading.runs[0]
            self.assertTrue(run.bold)
            self.assertAlmostEqual(run.font.size.pt, 12)
            self.assertEqual(run.font.color.rgb, document.RGBColor(0, 0, 0))
            self.assertAlmostEqual(heading.paragraph_format.space_before.pt, 8)
            self.assertAlmostEqual(heading.paragraph_format.space_after.pt, 4)
            self.assertTrue(heading.paragraph_format.keep_with_next)

        # Entry titles stay bold at the preset body size, not the heading size.
        title_paragraph = next(p for p in doc.paragraphs if p.text == "Acme Corp — Data Analyst")
        self.assertTrue(title_paragraph.runs[0].bold)
        self.assertAlmostEqual(title_paragraph.runs[0].font.size.pt, preset.body_font_pt)

    def test_contact_header_splits_phone_email_line_from_address(self):
        contact = {
            "Name": "Jordan Example",
            "Phone": "+1 555-0100",
            "Email": "jordan@example.com",
            "Address": "123 Main St, Springfield, IL",
        }
        build_docx(GOOD_DECISION, contact, ALL_CHUNKS, self.docx_path, FORMATTING_PRESETS[0])

        doc = DocxDocument(str(self.docx_path))
        texts = [p.text for p in doc.paragraphs]

        self.assertIn("+1 555-0100 | jordan@example.com", texts)
        self.assertIn("123 Main St, Springfield, IL", texts)
        # Address must not be merged onto the phone/email line.
        self.assertNotIn("+1 555-0100 | jordan@example.com | 123 Main St, Springfield, IL", texts)

        name_paragraph = next(p for p in doc.paragraphs if p.text == "Jordan Example")
        self.assertEqual(name_paragraph.alignment, document.WD_ALIGN_PARAGRAPH.LEFT)

    def test_raises_before_writing_when_job_header_cannot_be_extracted(self):
        broken_decision = DraftResult(
            professional_summary="Summary.",
            summary_evidence_ids=["job-3"],
            resume_sections=[
                ResumeSection(
                    section="job",
                    entries=[ResumeEntry(source_chunk_id="job-3", bullets=[Bullet(text="Did stuff.", evidence_ids=["job-3"])])],
                )
            ],
            selected_evidence=[EvidenceDecision(chunk_id="job-3", reason="x")],
            omitted_evidence=[],
        )

        with self.assertRaises(DocumentGenerationError) as ctx:
            build_docx(broken_decision, GOOD_CONTACT, [BAD_JOB_CHUNK], self.docx_path, FORMATTING_PRESETS[0])

        self.assertIn("job-3", str(ctx.exception))
        self.assertFalse(self.docx_path.exists())  # nothing partial was written


class TestRenderAndVerifyMocked(unittest.TestCase):
    """render_and_verify()'s own control flow, with convert_docx_to_pdf/
    count_pdf_pages mocked. This does not exercise real LibreOffice."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.docx_path = Path(self._tmp.name) / "resume.docx"

    def test_reports_unverified_when_soffice_missing(self):
        missing_soffice = Path(self._tmp.name) / "no_such_soffice.exe"

        result = render_and_verify(GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, soffice_path=missing_soffice)

        self.assertEqual(result.status, "unverified_no_libreoffice")
        self.assertTrue(self.docx_path.exists())
        self.assertIsNone(result.page_count)

    def test_stops_at_first_preset_that_hits_two_pages(self):
        with patch("document.convert_docx_to_pdf", return_value=Path("fake.pdf")) as mock_convert, patch(
            "document.count_pdf_pages", return_value=2
        ) as mock_count:
            result = render_and_verify(
                GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, soffice_path=Path(__file__)
            )

        self.assertEqual(result.status, "verified_two_pages")
        self.assertEqual(result.page_count, 2)
        self.assertEqual(mock_convert.call_count, 1)
        self.assertEqual(mock_count.call_count, 1)
        self.assertEqual(result.attempts, [("comfortable", 2)])

    def test_continues_to_tighter_presets_only_when_overlong(self):
        with patch("document.convert_docx_to_pdf", return_value=Path("fake.pdf")), patch(
            "document.count_pdf_pages", return_value=3
        ) as mock_count:
            result = render_and_verify(
                GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, soffice_path=Path(__file__)
            )

        self.assertEqual(result.status, "failed_wrong_page_count")
        self.assertEqual(result.page_count, 3)
        self.assertEqual(mock_count.call_count, len(FORMATTING_PRESETS))
        self.assertLessEqual(mock_count.call_count, 3)
        self.assertEqual(result.attempts, [("comfortable", 3), ("compact", 3), ("tight", 3)])
        self.assertNotIn("best", result.message.lower())

    def test_stops_immediately_on_undersized_render_without_trying_tighter_presets(self):
        """A render that comes back under two pages must stop right
        away - a tighter preset would only shrink it further, so none
        should be tried, and this must not be confused with success."""
        with patch("document.convert_docx_to_pdf", return_value=Path("fake.pdf")) as mock_convert, patch(
            "document.count_pdf_pages", return_value=1
        ) as mock_count:
            result = render_and_verify(
                GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, soffice_path=Path(__file__)
            )

        self.assertEqual(result.status, "failed_insufficient_content")
        self.assertEqual(result.page_count, 1)
        self.assertEqual(result.preset_name, "comfortable")
        self.assertEqual(mock_convert.call_count, 1)  # no tighter preset was tried
        self.assertEqual(mock_count.call_count, 1)
        self.assertEqual(result.attempts, [("comfortable", 1)])

    def test_each_attempt_gets_a_fresh_unique_directory(self):
        seen_dirs = []

        def _fake_convert(docx_path, attempt_dir, soffice_path, timeout_seconds=45):
            seen_dirs.append(attempt_dir)
            return Path("fake.pdf")

        with patch("document.convert_docx_to_pdf", side_effect=_fake_convert), patch(
            "document.count_pdf_pages", return_value=3
        ):
            render_and_verify(GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, soffice_path=Path(__file__))

        self.assertEqual(len(seen_dirs), len(FORMATTING_PRESETS))
        self.assertEqual(len(set(seen_dirs)), len(seen_dirs))  # every attempt directory is distinct
        for directory in seen_dirs:
            self.assertTrue(Path(directory).is_dir())  # each was actually created, not just a reused name

    def test_conversion_error_is_reported_without_retrying_further(self):
        with patch("document.convert_docx_to_pdf", side_effect=LibreOfficeConversionError("boom")) as mock_convert:
            result = render_and_verify(
                GOOD_DECISION, GOOD_CONTACT, ALL_CHUNKS, self.docx_path, soffice_path=Path(__file__)
            )

        self.assertEqual(result.status, "failed_conversion_error")
        self.assertEqual(mock_convert.call_count, 1)


class TestLoadSavedDraftAndRenderOnlyCli(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp_path = Path(self._tmp.name)

    def _write_saved_draft(self, filename: str, draft_raw: dict) -> Path:
        payload = {
            "schema_version": 1,
            "saved_at": "2026-01-01T00:00:00+00:00",
            "model": "gpt-test-2026-01-01",
            "usage": {"input_tokens": 1, "output_tokens": 2, "total_tokens": 3},
            "draft": draft_raw,
        }
        path = self.tmp_path / filename
        path.write_text(json.dumps(payload), encoding="utf-8")
        return path

    def test_load_saved_draft_round_trips_decision_to_dict(self):
        saved_path = self._write_saved_draft("good.json", decision_to_dict(GOOD_DECISION))

        loaded = load_saved_draft(saved_path)

        self.assertEqual(loaded["model"], "gpt-test-2026-01-01")
        self.assertEqual(loaded["draft"]["professional_summary"], GOOD_DECISION.professional_summary)

    def test_main_rejects_saved_draft_that_no_longer_matches_the_profile(self):
        """Evidence IDs that don't match the (mocked) current profile
        must fail validation - render_and_verify must never be called."""
        saved_path = self._write_saved_draft("stale.json", decision_to_dict(GOOD_DECISION))
        out_path = self.tmp_path / "resume.docx"

        with patch("document.load_profile", return_value=({}, {}, [GOOD_JOB_CHUNK])), patch(
            "document.render_and_verify"
        ) as mock_render:
            exit_code = document.main([str(saved_path), "--out", str(out_path)])

        self.assertEqual(exit_code, 1)
        mock_render.assert_not_called()

    def test_main_proceeds_to_render_when_draft_matches_profile(self):
        saved_path = self._write_saved_draft("matching.json", decision_to_dict(GOOD_DECISION))
        out_path = self.tmp_path / "resume.docx"

        fake_result = document.RenderResult(
            docx_path=out_path, pdf_path=None, page_count=2, preset_name="comfortable",
            status="verified_two_pages", message="ok",
        )
        with patch("document.load_profile", return_value=({}, GOOD_CONTACT, ALL_CHUNKS)), patch(
            "document.render_and_verify", return_value=fake_result
        ) as mock_render:
            exit_code = document.main([str(saved_path), "--out", str(out_path)])

        self.assertEqual(exit_code, 0)
        mock_render.assert_called_once()


if __name__ == "__main__":
    unittest.main()
